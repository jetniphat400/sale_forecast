"""Builds two tabs of index.html: the data of the Trend tab ("Trend Pricelist Omni 2024-2026") from a saved pull of cube_Sale_APD, and the executive summary tab
("สรุปผู้บริหาร", METRICS.md Sec.49: --pull-targets, --build-exec; see the section at the end of this file).

Two stages, both run by the monthly runner (src/monthly_refresh.py):
  --pull   one read-only database session: three aggregate queries over the tab's scope (rows grouped by item, day and status; the most frequent
           product name per item; one completeness query), saved under output/data/trend_pull/ (untracked).
  --build  reads the saved pull and the price list, computes the tab's data, checks it against the completeness query and writes it into index.html between
           the markers TREND-DATA-BEGIN and TREND-DATA-END. Nothing else of index.html is touched. Makes no database connection.

Scope (as built in August 2026, kept): Price List items of the visible sheets; cube_Sale_APD rows with revenue_type 'Omni Channel', status Actual and MPS
(kept apart as qty and sale columns), createDate from the configured start date (config trend_tab, source_table, status_basis, date_range). The time axis is
createDate (PO receipt date). Grain of the pull: item x calendar day x status; the months of the tab are summed from those daily rows, so the daily drill-down
and the months agree. The month of the pull date is incomplete: shown faded and left out of ADI and CV2 (Syntetos-Boylan cut-offs from config trend_tab).

The spec remark of every item (ok / conflict / nospec / nodata) follows the August 2026 logic (specs()); the database name shown in the tooltip is the item's
most frequent product name in the scope.
"""
import argparse
import html
import json
import logging
import math
import os
import re
import sys
from datetime import datetime

import pandas as pd
import yaml

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
logger = logging.getLogger("build_trend_tab")

BEGIN_MARK = "/* TREND-DATA-BEGIN */"
END_MARK = "/* TREND-DATA-END */"
CODE_PATTERN = re.compile(r"^[A-Za-z0-9._\-]+$")


class TrendTabError(Exception):
    """The Trend tab's data could not be built or failed its check; nothing is written."""


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Spec remark (August 2026 logic, used as given)
# ---------------------------------------------------------------------------------------------------------------------------------------------

TH = [('เควี', 'kv'), ('กิโลโวลท์', 'kv'), ('กิโลโวลต์', 'kv'), ('แอมป์', 'a'), ('แอมแปร์', 'a'), ('โวลท์', 'v'), ('โวลต์', 'v'), ('เฮิรตซ์', 'hz')]
UNIT = r'(?:kvar|kva|kv|ka|ma|amps?|a|var|va|v|hz|kw|w|mm2|mm|m)'
K = {'kv': ('v', 1000), 'ka': ('a', 1000), 'kva': ('va', 1000), 'kvar': ('var', 1000), 'kw': ('w', 1000), 'amp': ('a', 1), 'amps': ('a', 1), 'ma': ('a', 0.001)}


def specs(s: str) -> set:
    """The numeric specs written in a product description or name, as {'22000v', '630a', ...} (units folded to v, a, va, var, w; kV x 1000)."""
    s = s.lower()
    s = re.sub(r'(\d),(?=\d{3})', r'\1', s)
    s = s.replace('√3', '').replace('/3', '')
    for th, en in TH:
        s = s.replace(th, ' ' + en)
    s = re.sub(r'(\d)[\s\.]*(' + UNIT + r')(?![a-z0-9])', r'\1\2', s)
    out = set()
    for num, unit in re.findall(r'(\d+(?:\.\d+)?)(' + UNIT + r')(?![a-z0-9])', s):
        v = float(num)
        if unit in K:
            u, m = K[unit]
            v *= m
            unit = u
        out.add(f'{v:g}{unit}')
    return out


_UNIT_ORDER = ["v", "a", "va", "var", "w", "hz", "mm2", "mm", "m"]
_UNIT_LABEL = {"v": "V", "a": "A", "va": "VA", "var": "VAR", "w": "W", "hz": "HZ", "mm2": "mm2", "mm": "mm", "m": "m"}
_KILO = {"v": "kV", "a": "kA", "va": "kVA", "var": "kVAR", "w": "kW"}


def format_specs(values: set) -> str:
    """A spec set for the tooltip in plain string order, for example '22kV, 630A' (v, a, va, var and w from 1000 up are shown in k units)."""
    parsed = []
    for x in values:
        m = re.match(r'(\d+(?:\.\d+)?)([a-z0-9]+)$', x)
        parsed.append((_UNIT_ORDER.index(m.group(2)) if m.group(2) in _UNIT_ORDER else 99, float(m.group(1)), m.group(2)))
    out = []
    for _, num, unit in parsed:
        out.append(f"{num / 1000:g}{_KILO[unit]}" if unit in _KILO and num >= 1000 else f"{num:g}{_UNIT_LABEL.get(unit, unit)}")
    return ", ".join(sorted(out))


def spec_status(pl_descriptions: list, db_name, has_sales: bool) -> tuple:
    """(status, price-list specs, database specs): nodata when the item has no sales in the scope; nospec when either side has no spec; ok when the two are equal or
    one is contained in the other; conflict otherwise."""
    if not has_sales:
        return "nodata", set(), set()
    sp_pl = specs(" / ".join(pl_descriptions))
    sp_db = specs(db_name or "")
    if not sp_pl or not sp_db:
        return "nospec", sp_pl, sp_db
    if sp_pl == sp_db or sp_pl <= sp_db or sp_db <= sp_pl:
        return "ok", sp_pl, sp_db
    return "conflict", sp_pl, sp_db


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Demand classes
# ---------------------------------------------------------------------------------------------------------------------------------------------

def demand_stats(qty_complete: list) -> tuple:
    """(ADI, CV2) over the complete months: ADI = months / months with qty > 0; CV2 = population variance / mean^2 of qty over months with qty > 0. (None, None) when no month has qty > 0."""
    nz = [float(x) for x in qty_complete if x > 0]
    if not nz:
        return None, None
    adi = len(qty_complete) / len(nz)
    mean = sum(nz) / len(nz)
    var = sum((x - mean) ** 2 for x in nz) / len(nz)
    return adi, var / (mean * mean)


def demand_class(adi: float, cv2: float, adi_cut: float, cv2_cut: float) -> str:
    """Syntetos-Boylan: Smooth (ADI below, CV2 below), Erratic (ADI below, CV2 at or above), Intermittent (ADI at or above, CV2 below), Lumpy (both at or above)."""
    if adi < adi_cut:
        return "Smooth" if cv2 < cv2_cut else "Erratic"
    return "Intermittent" if cv2 < cv2_cut else "Lumpy"


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Configuration and the price list
# ---------------------------------------------------------------------------------------------------------------------------------------------

def load_config(root: str = PROJECT_ROOT) -> dict:
    with open(os.path.join(root, "config", "config.yaml"), encoding="utf-8") as f:
        return yaml.safe_load(f)


def pricelist_rows(root: str, config: dict) -> pd.DataFrame:
    """The visible-sheet rows of the price list with the division of the sheet (config sheet_to_division) and a cleaned description."""
    sys.path.insert(0, os.path.join(root, "src"))
    from pricelist_reader import load_visible_product_rows
    pl = load_visible_product_rows(os.path.join(root, "reference", "pricelist.xlsx"))
    to_div = config["sheet_to_division"]
    missing = set(pl["sheet"].unique()) - set(to_div)
    if missing:
        raise TrendTabError(f"sheets with product rows and no sheet_to_division entry: {sorted(missing)}")
    pl = pl.assign(division=pl["sheet"].map(to_div))
    for col in ("category", "type", "description"):
        pl[col] = pl[col].map(lambda v: " ".join(str(v).split()) if pd.notna(v) else "")
    return pl


def pricelist_label(root: str) -> str:
    """The quarter label of the price list (for example Q3'2026), read from the title cell of every visible product sheet; the sheets must agree."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(root, "reference", "pricelist.xlsx"), read_only=True, data_only=True)
    labels = set()
    for name in wb.sheetnames:
        ws = wb[name]
        if ws.sheet_state != "visible":
            continue
        rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
        if len(rows) < 3 or not any(h and "product code" in str(h).lower() for h in rows[2]):
            continue
        title = " ".join(str(v) for v in rows[0] if v)
        m = re.search(r"Price List\s+(\S+)", title)
        if not m:
            raise TrendTabError(f"price list sheet {name!r} has no 'Price List <label>' in its title row")
        labels.add(m.group(1))
    if len(labels) != 1:
        raise TrendTabError(f"the visible product sheets name {len(labels)} different price list labels: {sorted(labels)}")
    return labels.pop()


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Pull (one read-only session, aggregate queries only)
# ---------------------------------------------------------------------------------------------------------------------------------------------

def scope_where(codes: list, config: dict) -> str:
    """The WHERE clause of the tab's scope; item codes come from the price list and are checked against a strict pattern before they are written into SQL."""
    bad = [c for c in codes if not CODE_PATTERN.match(c)]
    if bad:
        raise TrendTabError(f"item codes with characters outside the allowed set: {bad[:5]}")
    t = config["trend_tab"]
    code_list = "','".join(sorted(codes))
    status_list = "','".join(config["status_basis"])
    return (f"revenue_type = '{t['revenue_type']}' AND status IN ('{status_list}') AND createDate >= '{config['date_range']['start']}' "
            f"AND itemcode IN ('{code_list}')")


def pull(root: str = PROJECT_ROOT, queries_extra=None) -> dict:
    """Runs the three queries in ONE session (no retry: a login failure raises), saves them under config trend_tab.pull_dir and returns them.
    `queries_extra(run_query)` may run further aggregate queries inside the same session (used once, by the task that proved the VAT treatment)."""
    import db
    config = load_config(root)
    pl = pricelist_rows(root, config)
    codes = sorted(pl["code"].unique())
    where = scope_where(codes, config)
    table = config["source_table"]
    out_dir = os.path.join(root, *config["trend_tab"]["pull_dir"].split("/"))
    os.makedirs(out_dir, exist_ok=True)
    extra = None
    with db.session():
        pulled_at = datetime.now().isoformat(timespec="seconds")
        daily = db.run_query(f"SELECT itemcode, CAST(createDate AS date) AS d, status, SUM(qty) AS q, SUM(sale) AS s, COUNT(*) AS n FROM {table} WHERE {where} "
                             f"GROUP BY itemcode, CAST(createDate AS date), status")
        names = db.run_query(f"SELECT itemcode, productName, COUNT(*) AS n FROM {table} WHERE {where} GROUP BY itemcode, productName")
        total = db.run_query(f"SELECT COUNT(*) AS n_rows, SUM(qty) AS q, SUM(sale) AS s, MIN(CAST(createDate AS date)) AS first_d, MAX(CAST(createDate AS date)) AS last_d "
                             f"FROM {table} WHERE {where}")
        if queries_extra is not None:
            extra = queries_extra(db.run_query, table, where)
    completeness = {"pulled_at": pulled_at, "n_rows": int(total["n_rows"].iloc[0]), "qty": float(total["q"].iloc[0]), "sale": float(total["s"].iloc[0]),
                    "first_date": str(total["first_d"].iloc[0])[:10], "last_date": str(total["last_d"].iloc[0])[:10], "n_codes_in_scope": len(codes)}
    daily["d"] = daily["d"].astype(str).str[:10]
    daily.to_csv(os.path.join(out_dir, "daily.csv"), index=False, encoding="utf-8")
    names.to_csv(os.path.join(out_dir, "names.csv"), index=False, encoding="utf-8")
    with open(os.path.join(out_dir, "completeness.json"), "w", encoding="utf-8") as f:       # written last: the runner reads its time as "pull finished"
        json.dump(completeness, f, ensure_ascii=False, indent=1)
    logger.info("Trend tab pull saved: %d daily rows, %d name rows, completeness %s", len(daily), len(names), completeness)
    return {"daily": daily, "names": names, "completeness": completeness, "extra": extra}


def pull_targets(root: str = PROJECT_ROOT, year: int = None) -> pd.DataFrame:
    """The revenue targets of one year from Cube_Target_PMIS, ONE read-only session (no retry: a login failure raises), AGGREGATE ONLY: year, division, revenue type, product type
    (and its category), the sum of TargetRevenueAmount, the number of rows and of null amounts. No person, customer or product code column is selected. Saved as targets.csv under
    config exec_summary.target_pull_dir (the year of the data month unless given) with the pull time in targets_meta.json (written last)."""
    import db
    config = load_config(root)
    ex = config["exec_summary"]
    if year is None:                                           # the year of the data month: the last month of the series the forecast uses
        monthly = pd.read_csv(os.path.join(root, "output", "data", "processed_all_divisions_monthly_qty.csv"), usecols=["year_month"])
        year = int(str(monthly["year_month"].max())[:4])
    out_dir = os.path.join(root, *ex["target_pull_dir"].split("/"))
    os.makedirs(out_dir, exist_ok=True)
    with db.session():
        pulled_at = datetime.now().isoformat(timespec="seconds")
        df = db.run_query(f"SELECT [Year] AS yr, [Division] AS division, [RevenueType] AS revenue_type, [ProoductCateName] AS category, [ProductTypeName] AS product_type, "
                          f"SUM([TargetRevenueAmount]) AS amount, COUNT(*) AS n_rows, SUM(CASE WHEN [TargetRevenueAmount] IS NULL THEN 1 ELSE 0 END) AS n_null "
                          f"FROM {ex['target_table']} WHERE [Year] = {year} GROUP BY [Year], [Division], [RevenueType], [ProoductCateName], [ProductTypeName]")
    df.to_csv(os.path.join(out_dir, "targets.csv"), index=False, encoding="utf-8")
    with open(os.path.join(out_dir, "targets_meta.json"), "w", encoding="utf-8") as f:
        json.dump({"pulled_at": pulled_at, "year": year, "n_rows_aggregated": int(df["n_rows"].sum()), "amount_all": float(df["amount"].fillna(0).sum())}, f, indent=1)
    logger.info("Executive summary targets saved: %d aggregate rows for %d", len(df), year)
    return df


def read_pull(root: str, config: dict) -> tuple:
    d = os.path.join(root, *config["trend_tab"]["pull_dir"].split("/"))
    paths = [os.path.join(d, n) for n in ("daily.csv", "names.csv", "completeness.json")]
    missing = [p for p in paths if not os.path.exists(p)]
    if missing:
        raise TrendTabError(f"the saved Trend pull is missing: {[os.path.basename(p) for p in missing]} (run build_trend_tab.py --pull)")
    daily = pd.read_csv(paths[0], dtype={"itemcode": str, "status": str}).fillna({"q": 0.0, "s": 0.0})
    names = pd.read_csv(paths[1], dtype={"itemcode": str}, keep_default_na=False)
    with open(paths[2], encoding="utf-8") as f:
        completeness = json.load(f)
    return daily, names, completeness


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------------------------------------------------------------------------

def month_range(first: str, last: str) -> list:
    return [str(p) for p in pd.period_range(first, last, freq="M")]


def top_names(names: pd.DataFrame) -> tuple:
    """({itemcode: most frequent product name}, number of items whose top count is shared by more than one name). A tie is broken by the name in alphabetical order."""
    best, ties = {}, 0
    for code, g in names.groupby("itemcode"):
        g = g.assign(productName=g["productName"].astype(str)).sort_values(["n", "productName"], ascending=[False, True], kind="mergesort")
        if len(g) > 1 and g["n"].iloc[0] == g["n"].iloc[1]:
            ties += 1
        best[code] = " ".join(g["productName"].iloc[0].split())
    return best, ties


def build_data(daily: pd.DataFrame, names: pd.DataFrame, completeness: dict, pl: pd.DataFrame, config: dict, pricelist_quarter: str) -> tuple:
    """(OMNI dict, MATCH dict, report dict) from the saved pull and the price list. Stops (TrendTabError) when the totals of the data differ from the completeness
    query, when a daily row falls outside the months, or when a day's rows do not sum to the item's months."""
    import reader_values as rv
    t = config["trend_tab"]
    pulled_at = pd.Timestamp(completeness["pulled_at"])
    first_month = str(config["date_range"]["start"])[:7]
    last_month = f"{pulled_at.year:04d}-{pulled_at.month:02d}"
    months = month_range(first_month, last_month)
    midx = {m: i for i, m in enumerate(months)}
    n_months = len(months)
    n_complete = n_months - 1                                   # the month of the pull date is incomplete
    daily = daily.copy()
    daily["ym"] = daily["d"].astype(str).str[:7]
    outside = daily[~daily["ym"].isin(midx)]
    if len(outside):
        raise TrendTabError(f"{len(outside)} daily rows fall outside the months {months[0]}..{months[-1]} (latest {outside['d'].max()})")
    codes = sorted(pl["code"].unique())
    best_name, n_ties = top_names(names)
    items, match = [], {}
    by_item = {c: g for c, g in daily.groupby("itemcode")}
    n_multi_sheet = int((pl.drop_duplicates(["sheet", "code"]).groupby("code")["sheet"].nunique() > 1).sum())
    n_multi_row = int((pl.groupby("code").size() > 1).sum())
    adi_cut, cv2_cut = float(t["sb_adi_cutoff"]), float(t["sb_cv2_cutoff"])
    dec = int(t["decimals"])
    tot_q = tot_s = tot_n = 0.0
    for code in codes:
        rows = pl[pl["code"] == code]
        first = rows.iloc[0]
        qa, qm, sa, sm = ([0.0] * n_months for _ in range(4))
        dd = {}
        g = by_item.get(code)
        n_rows_item = 0
        if g is not None:
            for r in g.itertuples():
                i = midx[r.ym]
                if r.status == "Actual":
                    qa[i] += float(r.q); sa[i] += float(r.s); k = (1, 2)
                else:
                    qm[i] += float(r.q); sm[i] += float(r.s); k = (3, 4)
                day = r.d[2:4] + r.d[5:7] + r.d[8:10]
                rec = dd.setdefault(int(day), [int(day), 0.0, 0.0, 0.0, 0.0])
                rec[k[0]] += float(r.q); rec[k[1]] += float(r.s)
                n_rows_item += int(r.n)
                tot_q += float(r.q); tot_s += float(r.s); tot_n += int(r.n)
        dd_list = [[d[0]] + [round(x, 4) for x in d[1:]] for _, d in sorted(dd.items())]
        qty_all = [a + b for a, b in zip(qa, qm)]
        has_sales = n_rows_item > 0
        adi, cv2 = demand_stats(qty_all[:n_complete])
        if sum(qty_all) == 0 and not has_sales:
            cls = "NoSale"
        elif adi is None:
            cls = "NoSale31M"                                    # sales only in the incomplete month: no ADI and CV2
        else:
            cls = demand_class(adi, cv2, adi_cut, cv2_cut)
        name = " / ".join(dict.fromkeys(rows["description"]))
        item = {"c": code, "n": name, "cat": first["category"] or "-", "pt": first["type"] or "-", "sh": sorted(rows["division"].unique()), "cls": cls,
                "adi": None if adi is None else round(adi, dec), "cv2": None if cv2 is None else round(cv2, dec),
                "qa": [round(x, 4) for x in qa], "qm": [round(x, 4) for x in qm], "sa": [round(x, 2) for x in sa], "sm": [round(x, 2) for x in sm], "dd": dd_list}
        if has_sales:
            item["dbn"] = best_name.get(code, "")
        if len(rows) > 1:
            item["rk"] = f"รหัสเดียวมี {len(rows)} รายการใน pricelist"
        # the daily rows of the item sum to its months, exactly (to the cent / to 4 decimals)
        dq = sum(d[1] + d[3] for d in dd_list); ds = sum(d[2] + d[4] for d in dd_list)
        if abs(dq - sum(qty_all)) > 1e-3 or abs(ds - (sum(sa) + sum(sm))) > 0.05:
            raise TrendTabError(f"{code}: the daily rows do not sum to the months")
        items.append(item)
        status, sp_pl, sp_db = spec_status(list(dict.fromkeys(rows["description"])), item.get("dbn"), has_sales)
        entry = {"s": status}
        if has_sales:
            entry["dbn"] = item["dbn"]
        if status == "conflict":
            entry["pl_spec"], entry["db_spec"] = format_specs(sp_pl), format_specs(sp_db)
        match[code] = entry
    # the tab's totals equal the completeness query, exactly (rows; qty to 4 decimals; sale to the cent)
    items_q = sum(sum(i["qa"]) + sum(i["qm"]) for i in items); items_s = sum(sum(i["sa"]) + sum(i["sm"]) for i in items)
    checks = {"rows": (int(tot_n), completeness["n_rows"]), "qty": (round(tot_q, 4), round(completeness["qty"], 4)), "sale": (round(tot_s, 2), round(completeness["sale"], 2)),
              "items_qty": (round(items_q, 3), round(completeness["qty"], 3)), "items_sale": (round(items_s, 1), round(completeness["sale"], 1))}
    bad = {k: v for k, v in checks.items() if v[0] != v[1]}
    if bad:
        raise TrendTabError(f"the Trend tab's totals differ from the completeness query: {bad}")
    omni = {"months": months, "n31": n_complete, "cates": sorted({i["cat"] for i in items}), "types": sorted({i["pt"] for i in items}), "items": items,
            "meta": {"pull": rv.thai_date_short(completeness["pulled_at"]), "pull_iso": completeness["pulled_at"][:10], "pricelist": pricelist_quarter,
                     "first_year": int(months[0][:4]), "last_year": int(months[-1][:4]),
                     "month_incomplete": rv.thai_month_short(months[-1]), "base_first": rv.thai_month_short(months[0]), "base_last": rv.thai_month_short(months[n_complete - 1])}}
    report = {"n_items": len(items), "n_months": n_months, "n_complete": n_complete, "months": [months[0], months[-1]], "n_codes_on_more_than_one_sheet": n_multi_sheet,
              "n_codes_with_more_than_one_row": n_multi_row, "n_name_ties": n_ties, "totals": {"rows": int(tot_n), "qty": tot_q, "sale": tot_s},
              "classes": pd.Series([i["cls"] for i in items]).value_counts().to_dict(), "spec": pd.Series([m["s"] for m in match.values()]).value_counts().to_dict(),
              "pull_date": completeness["pulled_at"]}
    return omni, match, report


def _json(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")


def write_index(index_path: str, omni: dict, match: dict) -> None:
    """Replaces the two data blocks of the Trend tab between the markers. The first run wraps the existing `const MATCH = ...;` and `const OMNI = ...;` lines in the markers."""
    with open(index_path, encoding="utf-8", newline="") as f:
        text = f.read()
    block = f"{BEGIN_MARK}\nconst MATCH = {_json(match)};\nconst OMNI = {_json(omni)};\n{END_MARK}"
    if BEGIN_MARK in text:
        if text.count(BEGIN_MARK) != 1 or text.count(END_MARK) != 1:
            raise TrendTabError("index.html holds the data markers more than once")
        a, b = text.index(BEGIN_MARK), text.index(END_MARK) + len(END_MARK)
        new = text[:a] + block + text[b:]
    else:
        m = re.search(r"^const MATCH = \{.*?\};\r?\nconst OMNI = \{.*?\};", text, re.S | re.M)
        if not m or text.count("const OMNI = ") != 1:
            raise TrendTabError("index.html: the existing MATCH and OMNI data blocks were not found where expected")
        new = text[:m.start()] + block + text[m.end():]
    with open(index_path, "w", encoding="utf-8", newline="") as f:
        f.write(new)


def run_build(root: str = PROJECT_ROOT, index_path: str = None) -> dict:
    """Builds the tab's data from the saved pull and writes it into index.html (or `index_path`). Returns the report of the build."""
    config = load_config(root)
    daily, names, completeness = read_pull(root, config)
    pl = pricelist_rows(root, config)
    omni, match, report = build_data(daily, names, completeness, pl, config, pricelist_label(root))
    path = index_path or os.path.join(root, config["trend_tab"]["index_file"])
    write_index(path, omni, match)
    report["written_to"] = os.path.relpath(path, root) if path.startswith(root) else path
    logger.info("Trend tab data written: %s", {k: v for k, v in report.items() if k != "totals"})
    return report


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Executive summary tab ("สรุปผู้บริหาร", METRICS.md Sec.49): figures computed at build time from the saved pulls, the forecast page's data and the recorded plan outputs.
# The tab's HTML is written into index.html between EXEC_BEGIN and EXEC_END. Makes no database connection (the targets come from pull_targets()).
# ---------------------------------------------------------------------------------------------------------------------------------------------

EXEC_BEGIN = "<!-- EXEC-TAB-BEGIN -->"
EXEC_END = "<!-- EXEC-TAB-END -->"


class ExecSummaryError(TrendTabError):
    """The executive summary could not be built or failed its gate; nothing is written."""


def exec_targets(root: str, config: dict, pl: pd.DataFrame, year: int) -> tuple:
    """({price list division: target in baht}, report) from the saved aggregate pull of Cube_Target_PMIS (config exec_summary): the Omni Channel revenue target of `year`, summed per
    division. Divisions whose database label is the Price List's own take their rows; the database's PEM103 rows are split between PEM103 and PEM107 by the Price List category of
    each row (a row whose category is on no sheet of the two, or on both, stays unallocated and is reported); other database divisions are listed in the report and not shown."""
    ex = config["exec_summary"]
    d = os.path.join(root, *ex["target_pull_dir"].split("/"))
    paths = [os.path.join(d, n) for n in ("targets.csv", "targets_meta.json")]
    if not all(os.path.exists(p) for p in paths):
        raise ExecSummaryError("the saved revenue-target pull is missing (run build_trend_tab.py --pull-targets)")
    with open(paths[1], encoding="utf-8") as f:
        meta = json.load(f)
    if int(meta["year"]) != int(year):
        raise ExecSummaryError(f"the saved revenue-target pull is for {meta['year']}, the data month is in {year}")
    df = pd.read_csv(paths[0])
    o = df[(df["revenue_type"] == ex["target_revenue_type"]) & (df["yr"] == year)]
    split = ex["target_division_split"]
    cat_div = {}
    for r in pl.itertuples():
        cat_div.setdefault(r.category, set()).add(r.division)
    by = {dv: 0.0 for dv in ex["target_division_same_name"] + list(split["to_by_type"])}
    unallocated, other = 0.0, {}
    for r in o.itertuples():
        amount = 0.0 if pd.isna(r.amount) else float(r.amount)
        div = str(r.division)
        if div in ex["target_division_same_name"]:
            by[div] += amount
        elif div == split["from"]:
            cands = cat_div.get(" ".join(str(r.category).split()), set()) & set(split["to_by_type"])
            if len(cands) == 1:
                by[next(iter(cands))] += amount
            else:
                unallocated += amount
        else:
            other[div] = other.get(div, 0.0) + amount
    has_row = set(o["division"].astype(str))
    for dv in list(by):                       # a division with no target row at all has no target (not a target of 0)
        if dv in ex["target_division_same_name"] and dv not in has_row:
            by[dv] = None
        elif dv in split["to_by_type"] and split["from"] not in has_row:
            by[dv] = None
    return by, {"pulled_at": meta["pulled_at"], "year": int(year), "omni_total": float(o["amount"].fillna(0).sum()), "unallocated": unallocated, "other_divisions": other}


def exec_month_units(forecast_units, booked_units) -> float:
    """Units counted for one item in one remaining month: the larger of the forecast units of the vintage and the units already on the books (Actual + MPS with forecast_date
    in the month). An item with no forecast row counts 0 forecast units (decision D2 of the user, 2026-10-09; METRICS.md Sec.49)."""
    return max(float(forecast_units or 0.0), float(booked_units or 0.0))


def exec_month_baht(forecast_units, booked_units, price) -> int:
    """Whole baht of exec_month_units x the item's unit price (rounded half up at item and month level, as the forecast page does); 0 for an item with no price."""
    if price is None:
        return 0
    return int(math.floor(exec_month_units(forecast_units, booked_units) * float(price) + 0.5))


def exec_booked_units(raw: pd.DataFrame, items: set, months: list) -> dict:
    """{(item, month): units} already on the books: Actual + MPS units of the forecast-scope `items` with forecast_date (not before createDate) in `months`, from the saved sales
    pull `raw` (columns itemcode, createDate, forecast_date, qty, status, revenue_type); Omni Channel only, the scope and key of the YTD."""
    d = raw[raw["revenue_type"].eq("Omni Channel") & raw["status"].isin(["Actual", "MPS"]) & raw["itemcode"].isin(items)].copy()
    d["createDate"] = pd.to_datetime(d["createDate"])
    d["forecast_date"] = pd.to_datetime(d["forecast_date"], errors="coerce")
    d = d[d["forecast_date"].notna() & (d["forecast_date"] >= d["createDate"])]
    d["ym"] = d["forecast_date"].dt.strftime("%Y-%m")
    g = d[d["ym"].isin(months)].groupby(["itemcode", "ym"])["qty"].sum()
    return {k: float(v) for k, v in g.items()}


def exec_watch(ts_rows: list, limit: float, plan_above: list, n_late: int, pem107_alert) -> list:
    """The lines of "เรื่องที่ต้องระวัง", each only while its condition holds: a group whose Tracking Signal is beyond plus or minus `limit` (below: forecast below the actual), each
    division-month of the plan above the historical peak, materials already late (n_late above 0), the PEM107 not-late alert (a dict with `since` and `pct`, or None)."""
    watch = []
    for r in ts_rows:
        ts = r["value"]
        if ts is not None and abs(ts) > limit:
            watch.append({"kind": "low" if ts < 0 else "high", "group": r["group"], "value": ts})
    for r in plan_above:
        watch.append({"kind": "plan", **r})
    if n_late > 0:
        watch.append({"kind": "material", "n": n_late})
    if pem107_alert:
        watch.append({"kind": "pem107", **pem107_alert})
    return watch


def exec_data(root: str = PROJECT_ROOT, plan_dir: str = None) -> dict:
    """Every figure of the tab (see METRICS.md Sec.49). Stops (ExecSummaryError) when the data month is not the last month of the vintage's fit window or the vintage does not cover a
    remaining month of the year."""
    sys.path.insert(0, os.path.join(root, "src"))
    import build_material_plan_page as bmp
    import build_operation_plan_page as bopp
    import build_report as br
    import maxmin_v1
    import operation_plan as op
    import reader_values as rv
    config = br.load_config()
    basis = config["report"]["price_basis"]
    vf = rv.vintage_facts(root)
    forward = br.gather_forward_forecast(vf, br.gather_scope_table())
    codes = sorted({it["item"] for ts in forward["divisions"].values() for t in ts for it in t["items"]})
    price_info = rv.unit_prices(codes, basis, vf["fit_first"], vf["fit_last"], os.path.join(root, "reference", "pricelist.xlsx"), root)
    baht = br.gather_baht(forward, price_info)
    sale = pd.read_csv(os.path.join(root, *basis["sale_file"].split("/")), usecols=["itemcode", "year_month", "sale", "division"])
    data_month = str(sale["year_month"].max())
    if data_month != vf["fit_last"]:
        raise ExecSummaryError(f"the last month of the saved series ({data_month}) is not the last month of the vintage's fit window ({vf['fit_last']})")
    year, mnum = int(data_month[:4]), int(data_month[5:7])
    remaining = [f"{year}-{m:02d}" for m in range(mnum + 1, 13)]
    missing = [m for m in remaining if m not in forward["months"]]
    if missing:
        raise ExecSummaryError(f"the vintage covers {forward['months'][0]} to {forward['months'][-1]}, not the remaining months {missing}")
    last_first = f"{year - 1}-01"
    if str(sale["year_month"].min()) > last_first:
        raise ExecSummaryError(f"the saved series starts {sale['year_month'].min()}, after {last_first} needed for last year's YTD")
    cur = sale[(sale["year_month"] >= f"{year}-01") & (sale["year_month"] <= data_month)].groupby("division")["sale"].sum()
    prev = sale[(sale["year_month"] >= last_first) & (sale["year_month"] <= f"{year - 1}-{mnum:02d}")].groupby("division")["sale"].sum()
    pl = pricelist_rows(root, load_config(root))
    targets, target_report = exec_targets(root, config, pl, year)
    primary = br.gather_primary_results()
    acc = br.gather_accuracy_vs_naive(config, primary)
    rel = {r["label"]: r["relative_mae"] for r in acc["rows"] if r["kind"] == "division"}
    th = acc["thresholds"]
    # the remaining months: per item the larger of the forecast units and the units already on the books (D2), whole baht at the item's unit price
    import forward_test_scoring as fts
    raw = pd.read_csv(fts.RAW_HISTORY_PATH, usecols=["itemcode", "createDate", "forecast_date", "qty", "status", "revenue_type"])
    booked = exec_booked_units(raw, set(codes), remaining)
    prices = price_info["prices"]
    booked_baht = {}
    rows = []
    for d in baht["divisions"]:
        new_by_month = {m: 0 for m in remaining}
        old_by_month = {m: baht["divisions"][d]["total"][forward["months"].index(m)] for m in remaining}
        for t in forward["divisions"][d]:
            for it in t["items"]:
                price = prices[it["item"]]["price"]
                for m in remaining:
                    b = exec_month_baht(0.0, booked.get((it["item"], m), 0.0), price)
                    booked_baht.setdefault(it["item"], {})[m] = b
                    new_by_month[m] += exec_month_baht(it["values"][forward["months"].index(m)], booked.get((it["item"], m), 0.0), price)
        fc_rem = float(sum(new_by_month.values()))
        ytd, ly = float(cur.get(d, 0.0)), float(prev.get(d, 0.0))
        proj = ytd + fc_rem
        tgt = targets.get(d)
        rows.append({"division": d, "ytd": ytd, "last_ytd": ly, "change": (ytd / ly - 1) if ly else None, "forecast_remaining": fc_rem, "projection": proj, "target": tgt,
                     "to_target": (proj / tgt) if tgt else None, "relative_mae": rel.get(d), "projection_page_only": ytd + float(sum(old_by_month.values())),
                     "by_month": {m: {"page": int(old_by_month[m]), "used": int(new_by_month[m])} for m in remaining}})
    tot = {k: sum(r[k] for r in rows) for k in ("ytd", "last_ytd", "forecast_remaining", "projection")}
    have = [r for r in rows if r["target"]]
    tot["target"] = sum(r["target"] for r in rows) if len(have) == len(rows) else None
    tot["change"] = (tot["ytd"] / tot["last_ytd"] - 1) if tot["last_ytd"] else None
    tot["to_target"] = (tot["projection"] / tot["target"]) if tot["target"] else None
    cfg_op = op.load_config(root)
    _im, dm, _meta = op.read_outputs(root, cfg_op, plan_dir)
    plan_above = [{"division": r.division, "month": rv.thai_month_short(r.month), "pct": bopp.fmt_pct(r.share_of_capacity)} for r in dm[dm["above_capacity"].astype(bool)].itertuples()]
    n_late = int(bmp.build_values(root, plan_dir)["n_late"])
    with open(os.path.join(root, "output", "summary", "task2b_part4_pem107_alert.json"), encoding="utf-8") as f:
        alert = json.load(f)
    watch = exec_watch([{"group": r["label"], "value": r["tracking_signal"]} for r in acc["rows"]], th["limit"], plan_above, n_late,
                       None if alert.get("not_late_from_may_2026_pct") is None else {"since": rv.thai_month_year(alert["split_date"]), "pct": f"{float(alert['not_late_from_may_2026_pct']):.1f}%"})
    pending = [r for r in maxmin_v1.pending_criteria_rows() if r["who"] == "ผู้บริหาร"]
    return {"data_month": data_month, "data_month_label": rv.thai_month_short(data_month), "year": year, "year_be": year + 543, "remaining_months": remaining, "booked_baht": booked_baht, "rows": rows, "total": tot,
            "n_forecast_divisions": len(rows), "n_beat_naive": sum(1 for r in rows if r["relative_mae"] is not None and r["relative_mae"] < th["pass"]), "thresholds": th,
            "watch": watch, "pending": pending, "target_report": target_report, "vintage_id": vf["vintage_id"]}


def _money(x) -> str:
    return f"{x / 1e6:,.1f}"


def _signed_pct(x) -> str:
    return f"{100 * x:+.1f}%"


def exec_render(data: dict, config: dict) -> str:
    """The tab's HTML (cards, table, the lists and the footer) from `data`; every text is the approved text of config exec_summary.text with its braces filled."""
    import reader_values as rv
    T = config["exec_summary"]["text"]
    esc = html.escape
    d, tot = data, data["total"]
    dash = T["dash"]
    month = d["data_month_label"]
    n_watch = len(d["watch"])
    cards = [
        (T["card1_label"].format(data_month=month), T["million_value"].format(value=_money(tot["ytd"])), T["card1_sub"].format(change=dash if tot["change"] is None else _signed_pct(tot["change"]))),
        (T["card2_label"].format(year=d["year_be"]), T["million_value"].format(value=_money(tot["projection"])),
         T["card2_sub"].format(pct=dash if tot["to_target"] is None else f"{100 * tot['to_target']:.1f}%")),
        (T["card3_label"], T["card3_value"].format(n=d["n_beat_naive"], N=d["n_forecast_divisions"]),
         T["card3_sub"].format(n_rounds=rv.backtest_rounds(config), decision_date=rv.thai_date_short(config["maxmin_v1"]["pending_criteria_values"]["decision_date"]))),
        (T["card4_label"], T["card4_value"].format(n=n_watch), T["card4_sub"]),
    ]
    cards_html = "".join(f'<div class="kpi" id="execCard{i}"><div class="k">{esc(a)}</div><div class="v">{esc(b)}</div><div class="s">{esc(c)}</div></div>' for i, (a, b, c) in enumerate(cards, 1))

    def cells(r, remark):
        change = dash if r["change"] is None else _signed_pct(r["change"])
        target = dash if r["target"] is None else _money(r["target"])
        to_target = dash if r["to_target"] is None else f"{100 * r['to_target']:.1f}%"
        rel = dash if r.get("relative_mae") is None else f"{r['relative_mae']:.2f}"
        return (f"<td>{_money(r['ytd'])}</td><td>{_money(r['last_ytd'])}</td><td>{change}</td><td>{_money(r['projection'])}</td><td>{target}</td><td>{to_target}</td>"
                f"<td>{rel}</td><td>{esc(remark)}</td>")

    body = []
    for r in d["rows"]:
        body.append(f'<tr data-division="{esc(r["division"])}"><td>{esc(r["division"])}</td>{cells(r, T["remark_no_target"] if r["target"] is None else "")}</tr>')
    body.append(f'<tr class="total-row" id="execTotalRow"><td>{esc(T["total_label"])}</td>{cells(dict(tot, relative_mae=None), "")}</tr>')
    for code in ("PEM104", "PEMC"):
        body.append(f'<tr data-division="{code}"><td>{code}</td>' + f"<td>{dash}</td>" * 7 + f"<td>{esc(T['remark_' + code])}</td></tr>")
    head = "".join(f"<th>{esc(c)}</th>" for c in T["columns"])
    watch_lines = []
    for w in d["watch"]:
        if w["kind"] in ("low", "high"):
            watch_lines.append(T["watch_" + w["kind"]].format(group=w["group"], value=f"{w['value']:.1f}"))
        elif w["kind"] == "plan":
            watch_lines.append(T["watch_plan"].format(division=w["division"], month=w["month"], pct=w["pct"]))
        elif w["kind"] == "material":
            watch_lines.append(T["watch_material"].format(n=w["n"]))
        else:
            watch_lines.append(T["watch_pem107"].format(since=w["since"], pct=w["pct"]))
    watch_html = "".join(f"<li>{esc(x)}</li>" for x in watch_lines)
    decide_html = "".join(f"<li>{esc(r['topic'])} · {esc(r['to_decide'])}</li>" for r in d["pending"])
    return (f'<div class="kpis" id="execCards">{cards_html}</div>'
            f'<p class="hint" id="execUnit">{esc(T["unit_note"])}</p>'
            f'<div class="table-scroll"><table id="execTable"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'
            f'<h3 id="execWatchTitle">{esc(T["watch_heading"])}</h3><ul id="execWatch">{watch_html}</ul>'
            f'<h3 id="execDecideTitle">{esc(T["decide_heading"])}</h3><ul id="execDecide">{decide_html}</ul>'
            f'<p><a id="execToAssump" href="#" onclick="omniShowTab(4);return false">{esc(T["decide_link"])}</a></p>'
            f'<p class="hint" id="execFooter">{esc(T["footer"].format(data_month=month))}</p>')


def exec_gate(data: dict, sales_report_path: str) -> dict:
    """The tab's checks: (1) the total row equals the sum of the division rows (YTD, last-year YTD, projection, target); (2) each division's projection equals its YTD plus, per item and
    remaining month, the larger of the baht on the rendered forecast page (its per-item baht rows) and the baht already on the books (data["booked_baht"]). Raises ExecSummaryError on any difference."""
    import reader_values as rv
    rows, tot = data["rows"], data["total"]
    for k in ("ytd", "last_ytd", "projection", "forecast_remaining"):
        if abs(tot[k] - sum(r[k] for r in rows)) > 1e-6:
            raise ExecSummaryError(f"the total row's {k} differs from the sum of the division rows")
    if tot["target"] is not None and abs(tot["target"] - sum(r["target"] for r in rows)) > 1e-6:
        raise ExecSummaryError("the total row's target differs from the sum of the division targets")
    with open(sales_report_path, encoding="utf-8") as f:
        text = f.read()
    labels = [rv.thai_month_short(x) for x in data["remaining_months"]]
    booked = data["booked_baht"]
    checked = 0
    for r in rows:
        m = re.search(r'id="fwd-baht-table-' + re.escape(r["division"]) + r'">.*?<thead><tr>(.*?)</tr></thead><tbody>(.*?)</tbody>', text, re.S)
        if not m:
            raise ExecSummaryError(f"the forecast page has no baht table for {r['division']}")
        heads = re.findall(r"<th>(.*?)</th>", m.group(1))[1:]
        if heads[:len(labels)] != labels:
            raise ExecSummaryError(f"the forecast page's first months {heads[:len(labels)]} are not the remaining months {labels}")
        total = 0
        for code, cells in re.findall(r'<tr class="fwd-bitem"[^>]*><td>(.*?)</td>((?:<td>[^<]*</td>)+)</tr>', m.group(2), re.S):
            code = html.unescape(re.sub(r"<[^>]+>", "", code.split(" <span")[0])).strip()
            vals = [0 if c == "-" else int(c.replace(",", "")) for c in re.findall(r"<td>(.*?)</td>", cells)]
            for i, mm in enumerate(data["remaining_months"]):
                total += max(vals[i], booked.get(code, {}).get(mm, 0))
            checked += 1
        if abs(r["projection"] - (r["ytd"] + total)) > 1e-6:
            raise ExecSummaryError(f"{r['division']}: the year-end projection {r['projection']} differs from YTD plus the larger of the forecast page's baht and the baht on the books per item and month {r['ytd'] + total}")
    return {"passed": True, "divisions": [r["division"] for r in rows], "months": labels, "items_checked": checked}


def write_between(path: str, begin: str, end: str, block: str) -> None:
    """Replaces the text between two markers of `path` by `block` (the markers stay); the markers must each be present exactly once."""
    with open(path, encoding="utf-8", newline="") as f:
        text = f.read()
    if text.count(begin) != 1 or text.count(end) != 1:
        raise TrendTabError(f"{os.path.basename(path)} does not hold the markers {begin} and {end} exactly once")
    a, b = text.index(begin) + len(begin), text.index(end)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(text[:a] + block + text[b:])


def run_exec_build(root: str = PROJECT_ROOT, index_path: str = None, sales_report_path: str = None, plan_dir: str = None) -> dict:
    """Builds the executive summary tab: data, gate, HTML written into index.html (or `index_path`) between the markers. Returns the build report (figures and gate)."""
    config = load_config(root)
    data = exec_data(root, plan_dir)
    gate = exec_gate(data, sales_report_path or os.path.join(root, "forecast", "sales_report.html"))
    write_between(index_path or os.path.join(root, config["trend_tab"]["index_file"]), EXEC_BEGIN, EXEC_END, "\n" + exec_render(data, config) + "\n")
    return {"data_month": data["data_month"], "year": data["year"], "total": data["total"], "n_watch": len(data["watch"]), "n_beat_naive": data["n_beat_naive"], "gate": gate,
            "target_report": data["target_report"], "vintage_id": data["vintage_id"]}


def main(argv=None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pull", action="store_true", help="one read-only database session; saves the pull")
    ap.add_argument("--build", action="store_true", help="builds the data from the saved pull into index.html")
    ap.add_argument("--pull-targets", action="store_true", help="one read-only session: the revenue targets of the year from Cube_Target_PMIS (aggregates only)")
    ap.add_argument("--build-exec", action="store_true", help="builds the executive summary tab of index.html from the saved pulls and the recorded outputs")
    a = ap.parse_args(argv)
    if not (a.pull or a.build or a.pull_targets or a.build_exec):
        ap.error("give --pull, --pull-targets, --build and/or --build-exec")
    try:
        if a.pull_targets:
            pull_targets()
        if a.pull:
            pull()
        if a.build:
            run_build()
        if a.build_exec:
            run_exec_build()
    except TrendTabError as e:
        logger.error("TREND TAB: %s", e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
