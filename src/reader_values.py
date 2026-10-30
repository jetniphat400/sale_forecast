"""Build-time values for reader text on forecast/sales_report.html and forecast/inventory.html.

CONVENTIONS.md "Dynamic values": every number a dashboard reader sees is computed at build time
from verified data, config or a verified recorded output file; none is typed into reader text.
Each function below names its source. A recorded output file must exist in the repository, and
its value was verified by an earlier task (cited in the function's docstring); a missing file
raises, it never falls back to a typed number.
"""
import os
import re

import pandas as pd

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SUMMARY_DIR = os.path.join(PROJECT_ROOT, "output", "summary")
SEGMENTATION_SCRIPT = os.path.join(PROJECT_ROOT, "src", "investigations", "task2b_part2_fulfilment_segmentation.py")


class ReaderValueError(Exception):
    """A value for reader text could not be computed from its source; the build stops."""


def _read_csv(name: str) -> pd.DataFrame:
    path = os.path.join(SUMMARY_DIR, name)
    if not os.path.exists(path):
        raise ReaderValueError(f"Recorded output file missing: output/summary/{name} -- refusing to "
                               f"render a typed number in its place.")
    return pd.read_csv(path)


def backtest_settings(config: dict) -> dict:
    """config.yaml backtest block (METRICS.md Sec.39): months and origin spacing."""
    b = config["backtest"]
    total = b["train_months"] + b["val_months"] + b["test_months"]
    return {"total_months": total, "holdout": b["holdout_months"],
            "min_train": b["min_train_months"], "step": b["origin_step_months"]}


def backtest_rounds(config: dict) -> int:
    """Number of rolling origins, from config.yaml's backtest block; same construction as
    src/backtest_rekeyed.get_origins (tests/test_reader_values.py checks the two agree)."""
    s = backtest_settings(config)
    last_train = s["total_months"] - s["holdout"]
    origins = list(range(s["min_train"], last_train + 1, s["step"]))
    if origins[-1] != last_train:
        origins.append(last_train)
    return len(origins)


def total_history_months(config: dict) -> int:
    """Months of history in the series (config.yaml backtest block: train + validation + test)."""
    return backtest_settings(config)["total_months"]


def count_no_current_minmax() -> tuple:
    """(items with no current Min/Max setting, items in the pilot scope) from
    output/summary/phaseE1fix_2_current_minmax.csv (column has_current_setting; 128 pilot items).
    Figure first stated in STATUS.md Phase E1-fix-2 and re-derived directly from this file."""
    df = _read_csv("phaseE1fix_2_current_minmax.csv")
    return int((~df["has_current_setting"].astype(bool)).sum()), int(len(df))


def current_minmax_date() -> str:
    """The date (ISO) the file behind count_no_current_minmax was written: the count is as of that day, not of today (the daily stock pull carries no Min and Max)."""
    import datetime as _dt
    return _dt.date.fromtimestamp(os.path.getmtime(os.path.join(SUMMARY_DIR, "phaseE1fix_2_current_minmax.csv"))).isoformat()


def forecast_date_disagreement_pct() -> float:
    """Share of rows where forecast_date disagrees with PlanDelDate, in percent: rows in
    output/summary/phaseA_a1_task2_plandeldate_disagreement_rows.csv over the full-scope row count
    in phaseA_a1_task2_join_match_summary.csv (n_apd_rows_forecast_date_notnull). Verified by the
    Phase A forecast_date task (docs/reports/summary/phaseA_a1_forecastdate_revision_findings.md,
    STATUS.md Phase A verification)."""
    rows = _read_csv("phaseA_a1_task2_plandeldate_disagreement_rows.csv")
    summ = _read_csv("phaseA_a1_task2_join_match_summary.csv")
    full = summ[summ["scope"] == "full_128_scope"]
    if full.empty:
        raise ReaderValueError("phaseA_a1_task2_join_match_summary.csv has no full_128_scope row")
    return 100.0 * len(rows) / float(full["n_apd_rows_forecast_date_notnull"].iloc[0])


def fulfilment_notice_threshold_days() -> int:
    """S3_THRESHOLD_DAYS of the segmentation that produced the item classes (METRICS.md Sec.23),
    read from src/investigations/task2b_part2_fulfilment_segmentation.py."""
    with open(SEGMENTATION_SCRIPT, encoding="utf-8") as f:
        m = re.search(r"^S3_THRESHOLD_DAYS\s*=\s*(\d+)", f.read(), re.M)
    if not m:
        raise ReaderValueError("S3_THRESHOLD_DAYS not found in the segmentation script")
    return int(m.group(1))


def fulfilment_signal_count() -> int:
    """Number of behaviour signals (S1, S2, S3, ...) in output/summary/task2b_part2_item_level.csv."""
    df = _read_csv("task2b_part2_item_level.csv")
    return len([c for c in df.columns if re.fullmatch(r"S\d+", c)])


def confirmed_lead_time_range(config: dict) -> tuple:
    """(first, last) days of the procurement lead time the business confirmed (STATUS.md Section 8.1): the smallest value of the grid and the default, which config marks as the upper end of the
    confirmed range. The last grid value (75) is a stress point only and is never part of it; a grid that does not extend beyond the default raises."""
    e1 = config["phase_e1_assumptions"]
    grid, default = e1["procurement_lead_time_days_grid"], e1["procurement_lead_time_days_default"]
    if not (min(grid) <= default < max(grid)):
        raise ReaderValueError("procurement_lead_time_days: the default must lie inside the grid and below its stress point")
    return min(grid), default


def not_forecast_reasons(config: dict = None) -> dict:
    """config `not_forecast_reasons`: {reason code: {divisions, line, min_max_line}}, the one entry that states why a division is not forecast."""
    import yaml
    if config is None:
        with open(os.path.join(PROJECT_ROOT, "config", "config.yaml"), encoding="utf-8") as f:
            config = yaml.safe_load(f)
    return config["not_forecast_reasons"]


def division_reason_codes(config: dict = None) -> dict:
    """{division: reason code} from config `not_forecast_reasons`; a division listed under two codes stops the build."""
    out = {}
    for code, entry in not_forecast_reasons(config).items():
        for division in entry["divisions"]:
            if division in out:
                raise ReaderValueError(f"{division} is listed under two not_forecast_reasons: {out[division]} and {code}")
            out[division] = code
    return out


def not_forecast_text(code: str, kind: str, division: str = None, config: dict = None):
    """The text of reason `code` of kind 'line' or 'min_max_line' with {division} filled, or None when config has no such reason (the caller then uses its general line)."""
    entry = not_forecast_reasons(config).get(code)
    if entry is None:
        return None
    return entry[kind].format(division=division)


def date_today():
    """Today's date on the machine's clock (one place, so a test can fix it)."""
    import datetime as _dt
    return _dt.date.today()


def _as_date(value):
    """A date from a date, a timestamp or a 'YYYY-MM-DD...' text."""
    import datetime as _dt
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    return _dt.date.fromisoformat(str(value).strip()[:10])


def previous_weekday(today, weekdays_back: int = 1):
    """The date `weekdays_back` Monday-to-Friday days before `today` (the day itself never counts): Monday gives the Friday before, Saturday and Sunday give Friday."""
    import datetime as _dt
    d, n = _as_date(today), 0
    while n < int(weekdays_back):
        d -= _dt.timedelta(days=1)
        if d.weekday() < 5:
            n += 1
    return d


def data_cutoff(today, config: dict):
    """The date a DATA source's date must not be before to read as fresh on `today`: the previous weekday (config page_timestamps.daily_weekdays_back weekdays back; holidays are not
    modelled)."""
    return previous_weekday(today, int(config["page_timestamps"]["daily_weekdays_back"]))


def source_class(config: dict, label: str) -> str:
    """'data', 'model_round' or 'computed': the class config page_timestamps gives the labelled source `label`. A source in none of the three lists, or in more than one, stops the build."""
    pt = config["page_timestamps"]
    found = [name for name, key in (("data", "data_sources"), ("model_round", "model_round_sources"), ("computed", "computed_sources")) if label in pt.get(key, [])]
    if len(found) != 1:
        raise ReaderValueError(f"the labelled source {label!r} is in {'more than one list' if found else 'none of the lists'} of config page_timestamps (data_sources, model_round_sources, computed_sources)")
    return found[0]


def data_is_stale(date, today, config: dict) -> bool:
    """True when a DATA source dated `date` (its pull date; for a static file its last data date) is old on `today`."""
    return _as_date(date) < data_cutoff(today, config)


def data_status_text(date, today, config: dict, file_date: bool = False) -> str:
    """The approved status wording for a DATA label on `today`. `file_date` is True for a static file that is not pulled (its stale text says its last data date, not a pull)."""
    t = config["page_timestamps"]["text"]
    return t["data_stale_file" if file_date else "data_stale"] if data_is_stale(date, today, config) else t["data_fresh"]


def data_age(date, shown: str, config: dict, today=None, file_date: bool = False) -> str:
    """HTML of a DATA label: `shown` and its status, in a span the page's script (DATA_AGE_JS) re-judges against the day the page is read; the text written here is the status on `today`
    (the build day by default), for a reader without the script and for the tests."""
    import html as _html
    status = data_status_text(date, today if today is not None else date_today(), config, file_date)
    return (f'<span class="data-age" data-pulled="{_html.escape(str(date)[:10])}"{" data-file=1" if file_date else ""}>{_html.escape(shown)} · <span class="data-status">{_html.escape(status)}</span></span>')


def data_age_js(config: dict) -> str:
    """The script that re-judges every `.data-age` element against the viewer's today (same rule as data_cutoff). Idempotent: a page that carries it twice runs it twice with the same result."""
    import json as _json
    t = config["page_timestamps"]["text"]
    params = _json.dumps({"back": int(config["page_timestamps"]["daily_weekdays_back"]), "fresh": t["data_fresh"], "stale": t["data_stale"], "stale_file": t["data_stale_file"]}, ensure_ascii=False)
    return ("(function(){var P=" + params + ";"
            "function iso(d){return d.getFullYear()+'-'+String(d.getMonth()+1).padStart(2,'0')+'-'+String(d.getDate()).padStart(2,'0');}"
            "function prev(back){var d=new Date(),n=0;while(n<back){d.setDate(d.getDate()-1);if(d.getDay()!==0&&d.getDay()!==6)n++;}return iso(d);}"
            "window.applyDataAge=function(){var cut=prev(P.back);document.querySelectorAll('.data-age').forEach(function(e){var s=e.querySelector('.data-status');"
            "if(s&&e.dataset.pulled)s.textContent=e.dataset.pulled<cut?(e.dataset.file?P.stale_file:P.stale):P.fresh;});};"
            "window.applyDataAge();document.addEventListener('DOMContentLoaded',window.applyDataAge);})();")


def forecast_header_text(config: dict, vf: dict = None) -> str:
    """The text of the "Forecast ล่าสุด" line, from the latest vintage's own metadata (`vintage_facts`: forecast_run_date of the log, fit_last_month of its metadata)."""
    vf = vf or vintage_facts()
    t = config["page_timestamps"]["text"]
    return t["header_forecast"].format(round=thai_month_short(vf["run_date"]), run=thai_date_short(vf["run_date"]), fit_last=thai_month_short(vf["fit_last"]))


def header_when(value) -> str:
    """A pull stamp as the header writes it: date and time when it has a time, else the date."""
    s = str(value).strip()
    return thai_datetime_short(s) if len(s) > 10 and s[10] in " T" else thai_date_short(s[:10])


def data_header(config: dict, items: list, forecast: bool = True, vf: dict = None, today=None) -> str:
    """The one standard header of a page or tab (Prompt 52): "ข้อมูลอัปเดตล่าสุด:" with every dated source the page uses, then "Forecast ล่าสุด:" where the page uses the forecast.
    `items` are (key of config page_timestamps.text.header_items, date or stamp, kind); kind 'data' (pulled: a status against the day the page is read), 'file' (a file's last data
    date: the file wording of the status), 'round' (what the forecast round was fitted on: a date and no status), 'computed' (a result calculated on that date from pulled data: a calculation
    date, no status, and the next monthly calculation `{next}`) or 'note' (a line with no date). The page build time is not part of it."""
    import html as _html
    t = config["page_timestamps"]["text"]
    rows = []
    for key, date, kind in items:
        if kind == "note":                       # a line with no date and no status (where to find today's figures)
            rows.append(_html.escape(t["header_items"][key]))
            continue
        shown = t["header_items"][key].format(when=header_when(date), next=next_calc_text(date, config) if kind == "computed" else "")
        rows.append(data_age(str(date)[:10], shown, config, today, file_date=(kind == "file")) if kind in ("data", "file") else _html.escape(shown))
    out = (f'<div class="data-header" id="data-header"><p class="note-line" id="data-header-data"><b>{_html.escape(t["header_data_label"])}</b> '
           f'{"<br>".join(rows) if rows else _html.escape(t["header_none"])}</p>')
    if forecast:
        out += f'<p class="note-line" id="data-header-forecast"><b>{_html.escape(t["header_forecast_label"])}</b> {_html.escape(forecast_header_text(config, vf))}</p>'
    return out + "</div>"


def model_round_facts(run_date: str) -> dict:
    """{"round": 'ต.ค. 69', "pulled": '5 ต.ค. 69', "pulled_iso": '2026-10-05'}: the run month of the forecast vintage (`run_date`, from its forward-test log rows) and the date the data of the round
    was pulled (the snapshot_pull_date of the series the report is built on, under this project's output/data)."""
    pulled = str(pd.read_csv(os.path.join(PROJECT_ROOT, "output", "data", "processed_full_category_sales_monthly_forecastDate.csv"), usecols=["snapshot_pull_date"])["snapshot_pull_date"].iloc[0])
    return {"round": thai_month_short(run_date), "pulled": thai_date_short(pulled), "pulled_iso": pulled[:10]}


def pipeline_gap_pct(pem101: dict, controls: dict, days_per_month: float) -> float:
    """Gap between the page's monthly-prorated stock value and the server-side daily-window
    figure for PEM101, in percent (absolute), at the default scenario, like for like: the items
    that the page gives a Min/Max AND that appear in output/summary/phaseE1fix_2_minmax_stockvalue.csv
    (column stock_value_contribution; Modeler/Validator agreement on its total recorded in
    STATUS.md Phase E1-fix-2, 0.09% apart)."""
    from inventory_recompute_reference import compute_all
    ref = _read_csv("phaseE1fix_2_minmax_stockvalue.csv").set_index("code")["stock_value_contribution"]
    page = compute_all(pem101, controls, days_per_month)
    rows = [p for p in page["per_item"] if p["stock_value_contribution"] > 0 and p["code"] in ref.index]
    page_total = sum(p["stock_value_contribution"] for p in rows)
    ref_total = float(ref.loc[[p["code"] for p in rows]].sum())
    if ref_total <= 0:
        raise ReaderValueError("recorded pipeline stock value is zero -- gap not computable")
    return abs(100.0 * (page_total - ref_total) / ref_total)


# Standard Thai month abbreviations (used for the month names of dates computed at build time). The year
# printed with them is the Gregorian year plus 543.
THAI_MONTH_ABBR = {1: "ม.ค.", 2: "ก.พ.", 3: "มี.ค.", 4: "เม.ย.", 5: "พ.ค.", 6: "มิ.ย.",
                   7: "ก.ค.", 8: "ส.ค.", 9: "ก.ย.", 10: "ต.ค.", 11: "พ.ย.", 12: "ธ.ค."}


def thai_month_year(iso_date: str) -> str:
    """'2026-05-01' -> 'พ.ค. 69': the one year format of the site (decision of the user, 2026-10-12): Thai short month and the Buddhist year as two digits. Same as thai_month_short."""
    return thai_month_short(iso_date)


def intermittent_lumpy_share_pct() -> float:
    """Share (percent) of forecast-status items whose demand is Intermittent or Lumpy, by METRICS.md
    Sec.29 (ADI = months / months with sales; CV2 = population variance of the non-zero months over
    their mean squared; Intermittent ADI >= 1.32 and CV2 < 0.49, Lumpy ADI >= 1.32 and CV2 >= 0.49).
    Computed from output/data/processed_all_divisions_monthly_qty.csv (Actual+MPS quantity, the 31
    months of the series) for the items with status_category == 'forecast' in
    output/summary/phaseC_step1revised_item_status_445.csv. Items with no sales in the window
    (no ADI) are left out of the denominator. The Sec.29 formula was checked item by item against the
    dashboard's own classification (0 mismatches, METRICS.md Sec.29)."""
    monthly = pd.read_csv(os.path.join(PROJECT_ROOT, "output", "data", "processed_all_divisions_monthly_qty.csv"))
    status = _read_csv("phaseC_step1revised_item_status_445.csv")
    forecast_codes = set(status[status["status_category"] == "forecast"]["itemcode"])
    n_classified = n_hit = 0
    for _code, g in monthly[monthly["itemcode"].isin(forecast_codes)].groupby("itemcode"):
        q = g["qty"].to_numpy(dtype=float)
        nonzero = q[q > 0]
        if len(nonzero) == 0:
            continue
        adi = len(q) / len(nonzero)
        cv2 = float(nonzero.var() / nonzero.mean() ** 2)
        n_classified += 1
        if adi >= 1.32:
            n_hit += 1        # Intermittent (CV2 < 0.49) or Lumpy (CV2 >= 0.49)
    if not n_classified:
        raise ReaderValueError("no forecast-status item has sales in the monthly series")
    return 100.0 * n_hit / n_classified


def first_scoring_month_label(config: dict) -> str:
    """Thai month and year of the first forward-test scoring round: the month of the first day after the
    first target month of output/summary/forward_test_log_all_divisions.csv has closed plus the leakage
    margin (config.yaml leakage_guard.min_margin_days), the same date step 6 of src/monthly_refresh.py
    reports as first_eligible_date."""
    log = pd.read_csv(os.path.join(SUMMARY_DIR, "forward_test_log_all_divisions.csv"), usecols=["target_month"])
    first_target = sorted(log["target_month"].unique())[0]
    window_end = pd.Period(first_target, freq="M").end_time.normalize()
    eligible = window_end + pd.Timedelta(int(config["leakage_guard"]["min_margin_days"]), unit="D")
    return thai_month_year(str((eligible + pd.Timedelta(1, unit="D")).date()))




# ---------------------------------------------------------------------------------------------------------------------------------
# One shared formatter for the dates the dashboard pages write (decision of the prompt of 2026-10-08): a round ('ต.ค. 69'), a month
# ('ส.ค. 69') and a date-time ('8 ต.ค. 69 08:01'), Buddhist year as two digits. The Min-Max page's script writes the same date-time form
# with the same month list (tests/test_inventory_page_browser.py checks the two agree); every Python page builder imports these.

def thai_month_short(ym: str) -> str:
    """'2026-10' (or any 'YYYY-MM...' text) -> 'ต.ค. 69'."""
    y, m = int(str(ym)[:4]), int(str(ym)[5:7])
    return f"{THAI_MONTH_ABBR[m]} {str(y + 543)[-2:]}"


def thai_datetime_short(s, seconds: bool = False) -> str:
    """'2026-10-06 08:00:40' -> '6 ต.ค. 69 08:00' (d MMM yy HH:mm, this machine's own clock, Buddhist year); seconds=True keeps the seconds where a page already shows them ('6 ต.ค. 69 08:00:40')."""
    t = pd.Timestamp(str(s)[:19])
    return f"{t.day} {THAI_MONTH_ABBR[t.month]} {str(t.year + 543)[-2:]} {t.strftime('%H:%M:%S' if seconds else '%H:%M')}"


def thai_year_short(year) -> str:
    """2026 -> '69' (the Buddhist year as two digits); a text writes it as 'ปี 69'."""
    return str(int(year) + 543)[-2:]


# The same formatter for the scripts of the pages: the Python builders put this text into their pages and index.html carries an identical copy between its two marker comments
# (tests/test_reader_text.py checks the copy and the output of both agree). No page has its own month list or year arithmetic.
JS_THAI_DATES = r"""/*TH-DATES-BEGIN*/
const THAI_MONTHS = ['ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.'];
function thaiYearShort(y) { return String(Number(y) + 543).slice(-2); }
function thaiMonthShort(ym) {
  const m = /^(\d{4})-(\d{2})/.exec(String(ym));
  return m ? THAI_MONTHS[Number(m[2]) - 1] + ' ' + thaiYearShort(m[1]) : String(ym);
}
function thaiDateShort(iso) {
  const m = /^(\d{4})-(\d{2})-(\d{2})/.exec(String(iso));
  return m ? Number(m[3]) + ' ' + THAI_MONTHS[Number(m[2]) - 1] + ' ' + thaiYearShort(m[1]) : String(iso);
}
function thaiDateTime(s, seconds) {
  const m = /^(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})(?::(\d{2}))?/.exec(s || '');
  if (!m) return 'ไม่ทราบ';
  return Number(m[3]) + ' ' + THAI_MONTHS[Number(m[2]) - 1] + ' ' + thaiYearShort(m[1]) + ' ' + m[4] + ':' + m[5] + (seconds && m[6] ? ':' + m[6] : '');
}
function thaiQuarter(label) { return String(label).replace(/Q(\d)'?(\d{4})/g, function (m, q, y) { return 'Q' + q + ' ' + thaiYearShort(y); }); }
// The one standard header of a page that loads its data in the browser (Prompt 52; the Python twin is reader_values.data_header). T = config page_timestamps.text with T.back (weekdays back);
// items = [{key, date, kind: 'data' | 'file' | 'round', id}]; forecastText = the "Forecast ล่าสุด" text or ''.
function dataHeaderHtml(T, items, forecastText) {
  const z = n => String(n).padStart(2, '0'), c = new Date(); let n = 0;
  while (n < T.back) { c.setDate(c.getDate() - 1); if (c.getDay() !== 0 && c.getDay() !== 6) n++; }
  const cut = c.getFullYear() + '-' + z(c.getMonth() + 1) + '-' + z(c.getDate());
  const esc = s => String(s).replace(/&/g, '&amp;').replace(/</g, '&lt;').replace(/>/g, '&gt;');
  const rows = items.map(function (it) {
    const s = String(it.date || '');
    let h = esc(T.header_items[it.key].replace('{when}', !s ? 'ไม่ทราบ' : (s.length > 10 ? thaiDateTime(s) : thaiDateShort(s))));
    if (it.kind !== 'round' && s) h += ' · <span class="data-status">' + esc(s.slice(0, 10) < cut ? (it.kind === 'file' ? T.data_stale_file : T.data_stale) : T.data_fresh) + '</span>';
    return '<span' + (it.id ? ' id="' + it.id + '"' : '') + '>' + h + '</span>';
  });
  return '<p class="note-line" id="data-header-data"><b>' + esc(T.header_data_label) + '</b> ' + (rows.length ? rows.join('<br>') : esc(T.header_none)) + '</p>' +
    (forecastText ? '<p class="note-line" id="data-header-forecast"><b>' + esc(T.header_forecast_label) + '</b> ' + esc(forecastText) + '</p>' : '');
}
/*TH-DATES-END*/"""


def next_calc_text(after, config: dict) -> str:
    """When the plans are next calculated after the stamp `after`, as "<date> <time>": the next weekday at the daily job's start (config page_timestamps.calc_times.daily), except on the monthly run
    day (`monthly_refresh.run_day`), which the monthly runner calculates at its own start (calc_times.monthly). Computed, never typed."""
    import datetime as _dt
    import monthly_refresh
    times = config["page_timestamps"]["calc_times"]
    text = str(after).strip()
    stamp = _dt.datetime.fromisoformat(text[:19].replace(" ", "T")) if len(text) > 10 else _dt.datetime.combine(_as_date(text), _dt.time())
    day = stamp.date()
    for _ in range(14):
        if day.weekday() < 5:
            hhmm = times["monthly"] if day == monthly_refresh.run_day(day.year, day.month) else times["daily"]
            if _dt.datetime.combine(day, _dt.time.fromisoformat(hhmm)) > stamp:
                return f"{thai_date_short(day.isoformat())} {hhmm}"
        day += _dt.timedelta(days=1)
    raise ReaderValueError(f"no calculation time found after {stamp}")


def criteria_decision_date(values: dict) -> str:
    """The date of the monthly run after which the draft criteria are decided: the run day (the first weekday on or after the 5th) of `values["decision_month"]` ('YYYY-MM'), as an ISO date."""
    import monthly_refresh
    year, month = str(values["decision_month"]).split("-")
    return monthly_refresh.run_day(int(year), int(month)).isoformat()


def thai_date_short(iso_date) -> str:
    """'2026-12-05' -> '5 ธ.ค. 69' (d MMM yy, Buddhist year as two digits)."""
    t = pd.Timestamp(str(iso_date)[:10])
    return f"{t.day} {THAI_MONTH_ABBR[t.month]} {str(t.year + 543)[-2:]}"


def vintage_facts(root: str = PROJECT_ROOT) -> dict:
    """What a page states about the latest forecast vintage it uses, read from the forward-test log and its metadata after the same hash check the plan
    uses (operation_plan.latest_vintage_forecast): vintage id, run date, the fit window (first and last month of the series the vintage was fitted on) and
    the forecast months, the forecast per item and month and the Item rows of the vintage (division and Type of each item). Stops when the metadata does not
    name the fit window."""
    import forward_test_common as ftc
    import operation_plan as op
    cfg = op.load_config(root)
    wide, vid, months, _ = op.latest_vintage_forecast(root, cfg)
    log = ftc.read_forward_test_log(op.path_of(root, cfg["forecast_log_file"]))
    meta = ftc.load_metadata(op.path_of(root, cfg["forecast_log_metadata_file"]))
    entry = meta[str(vid)] if str(vid) in meta else meta[vid]
    dates = sorted(set(log.loc[log["vintage_id"] == vid, "forecast_run_date"]))
    if len(dates) != 1:
        raise ReaderValueError(f"vintage {vid} has {len(dates)} run dates in the forward-test log: {dates}")
    for k in ("fit_first_month", "fit_last_month"):
        if not entry.get(k):
            raise ReaderValueError(f"the forward-test metadata of vintage {vid} has no {k}")
    item_rows = log[(log["vintage_id"] == vid) & (log["level"] == cfg["forecast_level"])][["itemcode", "division", "type", "target_month", "forecast_qty"]]
    return {"vintage_id": int(vid), "run_date": str(dates[0]), "fit_first": str(entry["fit_first_month"]), "fit_last": str(entry["fit_last_month"]),
            "forecast_months": [str(m) for m in months], "forecast": wide, "item_rows": item_rows}


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Unit price per item for the baht view of the forecast (decisions of the user, 2026-10-08; config.yaml report.price_basis holds the rule and the window).
# ---------------------------------------------------------------------------------------------------------------------------------------------

def market_prices(pricelist_path: str, basis: dict) -> tuple:
    """Market Price (labelled "Standard Price") of every product code, read from each visible product sheet of the price list.

    The column of each sheet is found by its header text (row 3 contains basis['market_price_header'], the row basis['market_price_label_row'] of the same column holds basis['market_price_label']),
    never by a column letter; a sheet with a product-code column and no such column, or with more than one, stops the build. Returns
    ({code: {'price': float or None, 'sheet': str, 'cell': 'P7'}}, {sheet: column letter}). A code on two sheets keeps its first sheet's row."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(pricelist_path, read_only=True, data_only=True)
    out, columns = {}, {}
    for name in wb.sheetnames:
        ws = wb[name]
        if ws.sheet_state != "visible":
            continue
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
        if len(rows) < 5:
            continue
        head2, head3 = rows[int(basis["market_price_label_row"]) - 1], rows[2]
        code_idx = next((i for i, h in enumerate(head3) if h and "product code" in str(h).lower()), None)
        if code_idx is None:
            continue
        hits = [i for i, h in enumerate(head3) if h and basis["market_price_header"].lower() in str(h).lower()
                and i < len(head2) and head2[i] and str(head2[i]).strip().lower() == basis["market_price_label"].lower()]
        if len(hits) != 1:
            raise ReaderValueError(f"price list sheet {name!r}: {len(hits)} columns match the Market Price header; exactly one is required")
        col = hits[0]
        columns[name] = get_column_letter(col + 1)
        for r_i, r in enumerate(rows[4:], start=5):
            code = r[code_idx] if code_idx < len(r) else None
            if code is None or str(code).strip() == "":
                continue
            code = str(code).strip()
            if code in out:
                continue
            v = r[col] if col < len(r) else None
            price = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
            out[code] = {"price": price, "sheet": name, "cell": f"{get_column_letter(col + 1)}{r_i}"}
    return out, columns


def price_window(fit_first: str, fit_last: str, available: list, window_months: int) -> tuple:
    """(recent months, fit-window months): the fit window of the vintage cut to the months the saved series holds, and its last `window_months` months."""
    fit = [m for m in available if fit_first <= m <= fit_last]
    if not fit or fit[-1] != fit_last:
        raise ReaderValueError(f"the saved monthly series does not reach the fit window {fit_first}..{fit_last} of the vintage the page uses")
    return fit[-window_months:], fit


def unit_prices(codes, basis: dict, fit_first: str, fit_last: str, pricelist_path: str, root: str = PROJECT_ROOT) -> dict:
    """Unit price per item code by the rule of config.yaml report.price_basis (tried in the order of basis['order']).

    Sale and qty are the saved monthly series (forecast-date key, Omni Channel, Price List items, Actual + MPS). Returns {'prices': {code: {'price', 'source'}},
    'recent': [months], 'fit': [months], 'market_columns': {sheet: column letter}}; source is one of the names in basis['order']; price is None only for 'none'."""
    sale = pd.read_csv(os.path.join(root, basis["sale_file"]), usecols=["itemcode", "year_month", "sale"])
    qty = pd.read_csv(os.path.join(root, basis["qty_file"]), usecols=["itemcode", "year_month", "qty"])
    both = sale.merge(qty, on=["itemcode", "year_month"], how="outer", validate="one_to_one")
    if both[["sale", "qty"]].isna().any().any():
        raise ReaderValueError("the saved sale and qty series do not cover the same item-months")
    recent, fit = price_window(fit_first, fit_last, sorted(both["year_month"].unique()), int(basis["window_months"]))
    sums = {}
    for key, months in (("recent", recent), ("fit", fit)):
        sums[key] = both[both["year_month"].isin(months)].groupby("itemcode")[["sale", "qty"]].sum()
    market, columns = market_prices(pricelist_path, basis)
    prices = {}
    for code in codes:
        chosen = None
        for rule in basis["order"]:
            if rule in ("sales_recent_window", "sales_fit_window"):
                g = sums["recent" if rule == "sales_recent_window" else "fit"]
                if code in g.index and g.at[code, "qty"] > 0:
                    chosen = {"price": float(g.at[code, "sale"] / g.at[code, "qty"]), "source": rule}
            elif rule == "market_price":
                m = market.get(code)
                if m and m["price"] is not None and m["price"] > 0:
                    chosen = {"price": m["price"], "source": rule}
            elif rule == "none":
                chosen = {"price": None, "source": rule}
            if chosen:
                break
        prices[code] = chosen
    return {"prices": prices, "recent": recent, "fit": fit, "market_columns": columns, "market": market}


def price_back_check(price_info: dict, basis: dict, root: str = PROJECT_ROOT) -> dict:
    """Sum(actual qty x unit price) / sum(actual sale) per division and calendar year of the saved monthly series, for the items that have a price.

    Inside the price window the ratio is 1 for items priced from that window; other years show how far the price has drifted. Returns {division: {year: ratio}};
    a ratio outside basis['back_check_low']..basis['back_check_high'] is listed in the key 'outside' as (division, year, ratio) so the caller can stop."""
    sale = pd.read_csv(os.path.join(root, basis["sale_file"]), usecols=["itemcode", "year_month", "sale", "division"])
    qty = pd.read_csv(os.path.join(root, basis["qty_file"]), usecols=["itemcode", "year_month", "qty"])
    d = sale.merge(qty, on=["itemcode", "year_month"], validate="one_to_one")
    d["price"] = d["itemcode"].map(lambda c: (price_info["prices"].get(c) or {}).get("price"))
    d["year"] = d["year_month"].str[:4]
    d = d[d["price"].notna()]
    d = d.assign(value=d["qty"] * d["price"])
    out, outside = {}, []
    for (div, year), g in d.groupby(["division", "year"]):
        if g["sale"].sum() == 0:
            continue
        ratio = float(g["value"].sum() / g["sale"].sum())
        out.setdefault(div, {})[year] = ratio
        if not basis["back_check_low"] <= ratio <= basis["back_check_high"]:
            outside.append((div, year, ratio))
    return {"ratios": out, "outside": outside}

# ---------------------------------------------------------------------------------------------------------------------------------------------------------------
# The numbers and dates of the user manual (docs/user_manual.md). The manual keeps {braced} names; the manual tab fills them at runtime from the key `manual_values` of
# data/assumptions.json (written with the assumptions by src/maxmin_v1.py). Each value below names its source; a missing source raises, nothing is typed.
# ---------------------------------------------------------------------------------------------------------------------------------------------------------------
def fulfilment_min_signals() -> int:
    """The number of signals that must hold for stock_policy (`n_hold >= N`), read from the segmentation script that produced the item classes (METRICS.md Sec.23)."""
    with open(SEGMENTATION_SCRIPT, encoding="utf-8") as f:
        m = re.search(r'"stock_policy" if n_hold >= (\d+)', f.read())
    if not m:
        raise ReaderValueError("the stock_policy signal rule (n_hold >= N) was not found in the segmentation script")
    return int(m.group(1))


def next_actual_info(data_month: str, run_date: str, guard_days: int, base_day: int = None) -> dict:
    """{"month": 'YYYY-MM', "run_date": ISO date}: the month after `data_month` (the last month of actual sales a page shows, the fit window's last month of the latest vintage) and the date of
    the monthly run that adds it. A month counts as complete `guard_days` after it ends (config leakage_guard.min_margin_days); the runs considered are the monthly run days
    (`monthly_refresh.run_day`) of the months after `run_date`, the run date of the vintage the page uses, so the answer changes only when a monthly run makes a new vintage. Computed, never typed."""
    import monthly_refresh
    y, m = (int(x) for x in str(run_date)[:7].split("-"))
    target = pd.Period(str(data_month), freq="M") + 1
    for k in range(1, 25):
        yy, mm0 = divmod(y * 12 + (m - 1) + k, 12)
        rd = monthly_refresh.run_day(yy, mm0 + 1, base_day)
        if (rd - target.end_time.date()).days >= guard_days:
            return {"month": str(target), "run_date": rd.isoformat()}
    raise ReaderValueError(f"no monthly run day within 24 months of {run_date} is {guard_days} days after the end of {target}")


def months_of_lag(guard_days: int, base_day: int, reference: str = "2026-10") -> int:
    """How many months the newest usable month trails the month of the run: the latest month whose end is at least `guard_days` before that month's run day (monthly_refresh.run_day)."""
    import monthly_refresh
    y, m = (int(x) for x in reference.split("-"))
    run = monthly_refresh.run_day(y, m, base_day)
    lag = 1
    while True:
        py, pm = (y, m - lag) if m - lag >= 1 else (y - 1, m - lag + 12)
        if (run - pd.Period(f"{py}-{pm:02d}", freq="M").end_time.date()).days >= guard_days:
            return lag
        lag += 1


# The manual_values read from sources the DAILY job refreshes (the Trend pull's completeness.json since Prompt 57; data/inventory.json with every daily stock publish). data/assumptions.json is written
# by the monthly run only, so its copy of these names is the value of the monthly build and differs from the live sources after any daily run: tests compare every other name exactly and check
# only the shape of these (Prompt 59).
DAILY_REFRESHED_MANUAL_VALUES = ("trend_data_through", "trend_first_month", "trend_last_month", "n_codes_trend", "n_codes_pages")


def manual_values(config: dict = None, root: str = PROJECT_ROOT) -> dict:
    """{name: text} for every {braced} name in docs/user_manual.md. Sources: config.yaml (page_timestamps, backtest, leakage_guard, shadow, phase_e1_assumptions, inventory_page), the Trend
    pull's completeness.json (what the Trend tab was built from), data/inventory.json (the codes the stock panel and the other pages cover), the notice-bucket recorded output and the
    segmentation script. Dates use the one site format: Thai short month and the Buddhist year as two digits."""
    import json
    import yaml
    if config is None:
        with open(os.path.join(root, "config", "config.yaml"), encoding="utf-8") as f:
            config = yaml.safe_load(f)
    with open(os.path.join(root, *config["trend_tab"]["pull_dir"].split("/"), "completeness.json"), encoding="utf-8") as f:
        trend = json.load(f)
    with open(os.path.join(root, "data", "inventory.json"), encoding="utf-8") as f:
        pages_codes = json.load(f)["totals"]["codes"]
    import backtest_all_divisions
    buckets = sorted(int(d) for d in _read_csv("leadtime_notice_buckets_overall.csv")["min_notice_days"])
    scenario = config["phase_e1_assumptions"]["default_scenario"]
    ranges = config["inventory_page"]["tier_a_ranges"]
    guard = int(config["leakage_guard"]["min_margin_days"])
    base = int(config["shadow"]["monthly_run_day_of_month"])
    return {
        "n_calibration_items": str(len(json.load(open(os.path.join(root, *config["maxmin_v1"]["ensemble_stock_items_file"].split("/")), encoding="utf-8")))),
        "stale_daily": str(int(config["page_timestamps"]["daily_weekdays_back"])),
        "n_rounds": str(backtest_rounds(config)), "holdout_months": str(config["backtest"]["holdout_months"]),
        "n_focus_items": str(len(backtest_all_divisions.FOCUS_ITEMS)),
        "notice_bucket_first": str(buckets[0]), "notice_bucket_second": str(buckets[1]), "notice_bucket_last": str(buckets[-1]),
        "trend_data_through": thai_date_short(trend["last_date"]), "trend_first_month": thai_month_short(trend["first_date"]), "trend_last_month": thai_month_short(trend["last_date"]),
        "n_codes_trend": str(trend["n_codes_in_scope"]), "n_codes_pages": str(pages_codes),
        "lead_confirmed_first": str(confirmed_lead_time_range(config)[0]), "lead_confirmed_last": str(confirmed_lead_time_range(config)[1]),
        "ctl_lead_time": str(scenario["procurement_lead_time_days"]), "ctl_assembly": str(scenario["assembly_time_days"]),
        "ctl_review": str(config["phase_e1_assumptions"]["review_interval_days_default"]), "ctl_service_level": f"{scenario['cycle_service_level']:.2f}",
        "ctl_holding_cost": f"{config['phase_e1_assumptions']['holding_cost_rate_annual']:.2f}",
        "csl_min": f"{ranges['cycle_service_level'][0]:.2f}", "csl_max": f"{ranges['cycle_service_level'][1]:.2f}",
        "notice_days": str(fulfilment_notice_threshold_days()), "signal_count": str(fulfilment_signal_count()), "min_signals": str(fulfilment_min_signals()),
        "run_day_base": str(base), "guard_days": str(guard), "lag_months": str(months_of_lag(guard, base)),
    }

